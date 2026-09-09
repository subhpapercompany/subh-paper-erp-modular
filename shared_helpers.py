# ==============================================================================
# SHARED HELPERS - SUBH PAPER ERP
# ==============================================================================
# Central helper functions, constants and DB access shared by every page module.
# Auto-generated from app_web_final_updated2.py; each page (tab) lives in its
# own page_*.py file so updates never delete unrelated code.
# ==============================================================================

import os
import re
import html
import sqlite3
import datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
import streamlit as st
import streamlit.components.v1 as components
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO

# ==============================================================================
# DATABASE CONFIGURATION & GLOBAL CONFIGS
# ==============================================================================
DB_NAME = "subh_paper_erp.db"
LOOKUP_FILE = "CAMLIN_CONVERSION_RATE_LOOKUP.xlsx"

MILL_OPTIONS = ["BGPPL", "TNPL", "Emami", "ITC", "Satia"]
LAMINATION_OPTIONS = ["Matt", "Gloss"]
DESTINATION_OPTIONS = ["Kolkata", "Guwahati", "Bengaluru"]
PARTY_OPTIONS = [
    "Emami Paper Mills Ltd.", "Satia Industries Ltd.", "Tamilnadu Newsprint Papers Ltd.",
    "Steller Product & Service", "Fortune Graphics", "Vishruta Packaging",
    "Salasar", "Rich", "Printask", "Our Godown"
]

# ==============================================================================
# DATE FORMATTING HELPERS - ALL DATES IN DD/MM/YYYY
# ==============================================================================
def format_date(date_value):
    """Convert date to dd/mm/yyyy format"""
    if date_value is None or date_value == '':
        return ''
    try:
        if isinstance(date_value, (datetime.date, datetime.datetime)):
            return date_value.strftime('%d/%m/%Y')
        elif isinstance(date_value, str):
            for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y'):
                try:
                    dt = datetime.datetime.strptime(date_value, fmt).date()
                    return dt.strftime('%d/%m/%Y')
                except ValueError:
                    continue
            return date_value
        else:
            return str(date_value)
    except Exception:
        return str(date_value)

def parse_date_input(date_str):
    """Parse dd/mm/yyyy input to datetime.date"""
    if not date_str:
        return None
    try:
        return datetime.datetime.strptime(date_str, '%d/%m/%Y').date()
    except ValueError:
        try:
            return datetime.datetime.strptime(date_str, '%Y-%m-%d').date()
        except ValueError:
            return None

def get_today_str():
    """Get today's date in dd/mm/yyyy format"""
    return datetime.date.today().strftime('%d/%m/%Y')

def get_date_input(label, key, default_value=None, help_text=None, use_calendar=True):
    """Date input helper. All date fields use the calendar picker by default."""
    if use_calendar:
        # Calendar picker with no automatic current-date selection.
        # The user must explicitly choose a date.
        calendar_value = None
        if isinstance(default_value, datetime.datetime):
            calendar_value = default_value.date()
        elif isinstance(default_value, datetime.date):
            calendar_value = default_value
        elif isinstance(default_value, str) and default_value:
            calendar_value = parse_date_input(default_value)

        selected_date = st.date_input(
            label,
            value=calendar_value,
            format="DD/MM/YYYY",
            key=key,
            help=help_text or "Select date from calendar"
        )

        if selected_date is None:
            return None, ""

        date_str = selected_date.strftime('%d/%m/%Y')
        return selected_date, date_str

    # All new Entry/Reporting screens must start with a clean date field.
    # A date is only shown when the caller explicitly supplies one.
    if default_value is None:
        default_value = ""
    elif isinstance(default_value, datetime.date):
        default_value = default_value.strftime('%d/%m/%Y')

    date_str = st.text_input(
        label,
        value=default_value,
        key=key,
        help=help_text or "Enter date in DD/MM/YYYY format"
    )

    parsed_date = parse_date_input(date_str)
    if parsed_date is None and date_str:
        st.warning(f"Invalid date format: {date_str}. Please use DD/MM/YYYY")
    return parsed_date, date_str

def get_month_year_options():
    current_year = datetime.datetime.now().year
    options = []
    for month in range(1, 13):
        month_name = datetime.date(current_year, month, 1).strftime("%B")
        options.append(f"{month_name} {current_year}")
    next_year = current_year + 1
    for month in range(1, 13):
        month_name = datetime.date(next_year, month, 1).strftime("%B")
        options.append(f"{month_name} {next_year}")
    return options

def get_month_only_options():
    """Order Month dropdown: sirf month names, April se March tak (FY order)."""
    return [datetime.date(2000, m, 1).strftime("%B") for m in [4,5,6,7,8,9,10,11,12,1,2,3]]

# ==============================================================================
# DATABASE FUNCTIONS
# ==============================================================================
def _read_env_file():
    """Read KEY=VALUE pairs from the .env file next to this module."""
    env_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env")
    values = {}
    if os.path.isfile(env_path):
        try:
            with open(env_path, encoding="utf-8") as _fh:
                for _line in _fh:
                    _line = _line.strip()
                    if not _line or _line.startswith("#") or "=" not in _line:
                        continue
                    _k, _, _v = _line.partition("=")
                    values[_k.strip()] = _v.strip().strip('"').strip("'")
        except Exception:
            pass
    return values

def _turso_credentials():
    env = _read_env_file()
    url = os.environ.get("TURSO_URL") or env.get("TURSO_URL")
    token = os.environ.get("TURSO_AUTH_TOKEN") or env.get("TURSO_AUTH_TOKEN")
    return url, token

def get_db_connection(private=False):
    """Open a database connection.

    If Turso credentials are present (env vars or .env), connect to the Turso
    cloud database so everyone sees the same data. Otherwise fall back to the
    local SQLite file (used for development / offline).

    private=True returns a fresh connection the caller owns and MUST close.
    private=False returns a session-cached connection that is reused across
    reruns and must NEVER be closed by the caller (it is process/session-owned).
    """
    url, token = _turso_credentials()
    if url and token:
        try:
            import libsql
            if private:
                return libsql.connect(url, auth_token=token, autocommit=True, timeout=30)
            # Reuse one connection per Streamlit session: har rerun me naya
            # handshake/connection banane se reruns bhut slow hote hain aur
            # click events race/lost ho jate hain. autocommit=True har statement
            # ko apne txn me rakhta hai, is liye idle session pe SQLITE_BUSY nahi aata.
            try:
                conn = st.session_state.get("_turso_conn")
            except Exception:
                conn = None
            if conn is None:
                conn = libsql.connect(url, auth_token=token, autocommit=True, timeout=30)
                try:
                    st.session_state["_turso_conn"] = conn
                except Exception:
                    pass
            return conn
        except Exception:
            pass
    return sqlite3.connect(DB_NAME)

@st.cache_data(ttl=8, show_spinner=False)
def _cf_item_options_cached():
    """CF item list (Excel Sheet3 codes + saved consumable entries), cached."""
    items = []

    # Product-code based CFC items from the master Excel file.
    try:
        cf_lookup_files = [
            "master_data/case.xlsx",
            "case.xlsx",
        ]
        cf_lookup_path = next((p for p in cf_lookup_files if os.path.exists(p)), None)
        if cf_lookup_path:
            cf_master_df = pd.read_excel(cf_lookup_path, sheet_name="Sheet3", header=1)
            if "Product Code" in cf_master_df.columns:
                for code in cf_master_df["Product Code"].dropna().tolist():
                    if isinstance(code, float) and code.is_integer():
                        code = int(code)
                    code = str(code).strip()
                    if code:
                        item_name = f"{code} CFC"
                        if item_name not in items:
                            items.append(item_name)
    except Exception:
        pass

    # Manually created items are kept in the database and remain available
    # in both Entry Mode and Reporting Mode.
    try:
        _c = get_db_connection(private=True)
        try:
            rows = _c.execute(
                "SELECT DISTINCT item_name FROM consumable_cf_entries "
                "WHERE item_name IS NOT NULL AND TRIM(item_name)<>'' "
                "ORDER BY item_name"
            ).fetchall()
            for _row in rows:
                item_name = str(_row[0]).strip()
                if item_name and item_name not in items:
                    items.append(item_name)
        finally:
            try:
                _c.close()
            except Exception:
                pass
    except Exception:
        pass

    return items

def get_consumable_cf_item_options(conn):
    return _cf_item_options_cached()

def fetch_case_mrp(product_code):
    """Return the Revised MRP from case.xlsx (Sheet3) for the matching product code."""
    code = str(product_code or "").strip()
    if not code:
        return None
    try:
        _here = os.path.dirname(os.path.abspath(__file__))
        case_lookup_files = [
            "master_data/case.xlsx",
            "case.xlsx",
            os.path.join(os.path.dirname(_here), "case.xlsx"),
            os.path.join(_here, "..", "case.xlsx"),
        ]
        case_path = next((p for p in case_lookup_files if os.path.exists(p)), None)
        if not case_path:
            return None
        case_df = pd.read_excel(case_path, sheet_name="Sheet3", header=1)
        if "Product Code" not in case_df.columns or "Revised MRP" not in case_df.columns:
            return None
        match = case_df[case_df["Product Code"].astype(str).str.strip() == code]
        if match.empty:
            return None
        value = match.iloc[0]["Revised MRP"]
        if value is None or str(value).strip() == "":
            return None
        return float(value)
    except Exception:
        return None

def parse_product_code(p_code):
    if not p_code or len(p_code) < 10:
        return None, None, None
    try:
        page = int(str(p_code)[-3:])
    except (ValueError, TypeError):
        page = None
    last_4_first = str(p_code)[-4:-3] if len(p_code) >= 4 else ""
    ruling_map = {"1": "Unruled", "3": "Single Line", "5": "Four Line", "8": "Med. Square"}
    ruling_type = ruling_map.get(last_4_first, "")
    third_char = str(p_code)[2:3] if len(p_code) >= 3 else ""
    return page, ruling_type, third_char

def _find_making_rate_chart_file():
    """Locate the user supplied Making rate chart workbook."""
    _here = os.path.dirname(os.path.abspath(__file__))
    _parent = os.path.dirname(_here)
    candidates = [
        "Making rate chart.xlsx",
        "MAKING RATE CHART.xlsx",
        "Making Rate Chart.xlsx",
        os.path.join(_here, "Making rate chart.xlsx"),
        os.path.join(_here, "MAKING RATE CHART.xlsx"),
        os.path.join(_parent, "MAKING RATE CHART.xlsx"),
        os.path.join(_parent, "Making Rate Chart.xlsx"),
        os.path.join(_parent, "Making rate chart.xlsx"),
        os.path.join(os.getcwd(), "Making rate chart.xlsx"),
    ]
    for path in candidates:
        if os.path.isfile(path):
            return path
    return ""

def fetch_making_rate(product_code):
    """
    Fetch Making Rate from Making Rate Chart.xlsx.
    Rule:
      - Product Code 3rd digit = 3 -> D/C chart, matching Page Count.
      - Product Code 3rd digit = 5 -> A4 chart, matching Page Count.
    """
    code = str(product_code or "").strip()
    if len(code) < 3:
        return 0.0

    third_digit = code[2]
    if third_digit not in {"3", "5"}:
        return 0.0

    try:
        page = int(code[-3:])
    except (TypeError, ValueError):
        return 0.0

    path = _find_making_rate_chart_file()
    if not path:
        return 0.0

    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        target_chart = "D/C" if third_digit == "3" else "A4"
        active_chart = ""

        for ws in wb.worksheets:
            active_chart = ""
            for r in range(1, ws.max_row + 1):
                b_value = str(ws.cell(r, 2).value or "").strip().upper()

                if b_value == target_chart:
                    active_chart = target_chart
                    continue

                if active_chart != target_chart:
                    continue

                page_value = ws.cell(r, 2).value
                rate_value = ws.cell(r, 3).value
                page_text = str(page_value or "").strip().upper()
                match = re.fullmatch(r"(\d+)\s*PG", page_text)

                if not match:
                    continue

                if int(match.group(1)) != page:
                    continue

                try:
                    return float(rate_value or 0)
                except (TypeError, ValueError):
                    rate_match = re.search(r"[-+]?\d*\.?\d+", str(rate_value or ""))
                    return float(rate_match.group()) if rate_match else 0.0

        return 0.0
    except Exception:
        return 0.0


def fetch_po_request_case_qty(conn, product_code, order_month=None):
    """Return PO Request Report column 10 (Case Qty) for the matching product."""
    code = str(product_code or "").strip()
    if not code:
        return 0.0
    try:
        if order_month:
            q = pd.read_sql_query(
                "SELECT * FROM po_request_entries WHERE product_code = ? AND po_month = ? ORDER BY id DESC LIMIT 1",
                conn, params=(code, order_month)
            )
        else:
            q = pd.read_sql_query(
                "SELECT * FROM po_request_entries WHERE product_code = ? ORDER BY id DESC LIMIT 1",
                conn, params=(code,)
            )
        if q.empty and order_month:
            q = pd.read_sql_query(
                "SELECT * FROM po_request_entries WHERE product_code = ? ORDER BY id DESC LIMIT 1",
                conn, params=(code,)
            )
        if q.empty:
            return float(_case_qty_from_product_code(code))
        report_row = calculate_po_request_full_row(q.iloc[0].to_dict())
        return float(report_row[9] or 0)
    except Exception:
        return float(_case_qty_from_product_code(code))

def _case_qty_from_product_code(product_code):
    """Compute PO Request Report Column 10 (Case Qty) directly from Product Code."""
    code = str(product_code or "").strip()
    if len(code) < 3:
        return 0.0
    fixed_120_codes = {
        "1132101132", "1132201140", "1132103132",
        "1132203140", "1132205140", "1132208140",
    }
    try:
        page = int(code[-3:]) if code[-3:].isdigit() else 0
    except (ValueError, TypeError):
        page = 0
    if code in fixed_120_codes:
        page = 120
    third_char = code[2:3]
    if third_char == "5":
        return float({52:192, 76:144, 100:120, 140:84, 176:72, 240:48}.get(page, 0))
    if third_char == "3":
        return float({56:288, 92:192, 120:144, 164:96, 176:96}.get(page, 0))
    return 0.0

def fetch_production_tab5_calculations(conn, product_code, production_qty):
    """
    Production Entry auto-calculations copied from GUI Tab 5 logic.

    Returns: (book_weight, paper_consumption, board_size, board_consumption)
    """
    code = str(product_code or "").strip()
    try:
        qty = float(production_qty or 0)
    except (TypeError, ValueError):
        qty = 0.0

    if len(code) < 3 or code[2] not in ("3", "5"):
        return 0.0, 0.0, get_board_size_from_product(code), 0.0

    third_digit = code[2:3]
    col6 = None
    col13 = col14 = col15 = col16 = col17 = col18 = None

    try:
        # Same first source used by GUI Tab 5: Production/Tab-1 data.
        source = pd.read_sql_query(
            "SELECT * FROM production_entries WHERE product_code = ? ORDER BY id DESC LIMIT 1",
            conn, params=(code,)
        )
        if not source.empty:
            full_row = calculate_po_request_full_row(source.iloc[0].to_dict())
            col6 = float(full_row[5])
            col13 = float(full_row[12])
            col14 = float(full_row[13])
            col15 = float(full_row[14])
            col16 = float(full_row[15])
            col17 = float(full_row[16])
            col18 = float(full_row[17])
        else:
            # Same fallback used by GUI Tab 5: Order Entry data.
            order = pd.read_sql_query(
                "SELECT order_month, product_code, ruling_type, "
                "no_of_page AS page, product_size AS book_size, order_qty "
                "FROM order_entries WHERE product_code = ? ORDER BY id DESC LIMIT 1",
                conn, params=(code,)
            )
            if not order.empty:
                try:
                    col6 = int(order.iloc[0]["page"])
                except (ValueError, TypeError):
                    col6 = 0

                if third_digit == "5":
                    col13, col14, col15 = 190, 45.5, 91
                    col16, col17, col18 = 57, 90, 43
                else:
                    col13, col14, col15 = 190, 98, 77
                    col16, col17, col18 = 57, 97, 37
    except (ValueError, TypeError, IndexError, sqlite3.Error, Exception):
        pass

    # If no database source exists yet, use the same product-code dimensions
    # used by the GUI fallback so the formula can still work.
    if col6 is None:
        try:
            col6 = int(code[-3:])
        except (ValueError, TypeError):
            col6 = 0

    if col13 is None:
        if third_digit == "5":
            col13, col14, col15 = 190, 45.5, 91
            col16, col17, col18 = 57, 90, 43
        else:
            col13, col14, col15 = 190, 98, 77
            col16, col17, col18 = 57, 97, 37

    try:
        base_weight = (col16 * col17 * col18) / 20000 / 493.75
        # Production Entry: Book Weight is always divided by 500 (per user rule),
        # not by 16 or 12.
        if third_digit == "3":
            book_weight_value = base_weight * (col6 / 500)
            board_consumption_value = qty / 8
        else:
            book_weight_value = base_weight * (col6 / 500)
            board_consumption_value = qty / 3

        paper_consumption_value = book_weight_value * qty
        board_size = f"{int(col15)}x{int(col14)}x{int(col13)}"

        # GUI Tab 5 uses round() for paper and board consumption.
        paper_consumption = float(round(paper_consumption_value))
        board_consumption = float(round(board_consumption_value))
        return float(book_weight_value), paper_consumption, board_size, board_consumption
    except (ValueError, TypeError, ZeroDivisionError):
        return 0.0, 0.0, get_board_size_from_product(code), 0.0


def get_product_size_from_code(p_code):
    if not p_code or len(p_code) < 3:
        return "24x18"
    return "29.7x21" if p_code[2:3] == "5" else "24x18"

def get_board_size_from_product(p_code):
    if not p_code or len(p_code) < 3:
        return "77x98x190"
    return "91x91x190" if p_code[2:3] == "5" else "77x98x190"

def fetch_purchase_order_product_size(conn, product_code, order_month=None):
    """Return the Product Size stored in the Purchase Order (order_entries).

    Prefers the row for the matching Order Month, then falls back to the most
    recent Purchase Order row for the Product Code.
    """
    code = str(product_code or "").strip()
    if not code:
        return ""
    try:
        po = pd.read_sql_query(
            """SELECT product_size FROM order_entries
               WHERE CAST(product_code AS TEXT) = ?
                 AND (? IS NULL OR ? = '' OR order_month = ?)
               ORDER BY id DESC LIMIT 1""",
            conn, params=(code, order_month, order_month, order_month)
        )
        if po.empty:
            po = pd.read_sql_query(
                "SELECT product_size FROM order_entries WHERE CAST(product_code AS TEXT) = ? ORDER BY id DESC LIMIT 1",
                conn, params=(code,)
            )
        if not po.empty:
            size = str(po.iloc[0].get("product_size", "") or "").strip()
            if size:
                return size
    except Exception:
        pass
    return ""

def _normalise_lookup_text(value):
    value = str(value or "").strip().upper()
    value = value.replace("\n", " " ).replace("\r", " " )
    value = re.sub(r"\s+", " ", value)
    compact = re.sub(r"[^A-Z0-9]+", "", value)
    if compact in ("MATT", "MATTE"):
        return "MATTE"
    if compact in ("GLOSS", "GLOSSY"):
        return "GLOSS"
    if compact in ("SINGLELINE", "SINGLELINES"):
        return "SINGLE LINE"
    if compact in ("FOURLINE", "FOURLINES"):
        return "FOUR LINE"
    if compact in ("MEDSQUARE", "MEDIUMSQUARE", "MEDSQUARES"):
        return "MED. SQUARE"
    if compact in ("UNRULED", "UNRULE"):
        return "UNRULED"
    return value

def _lookup_number(value):
    try:
        return int(float(str(value).strip()))
    except (ValueError, TypeError):
        return None

def _find_conversion_lookup_file():
    candidates = [
        LOOKUP_FILE,
        os.path.join(os.path.dirname(os.path.abspath(__file__)), LOOKUP_FILE),
        os.path.join(os.getcwd(), LOOKUP_FILE),
    ]
    # Also walk up parent directories (the workbook may sit one or more
    # folders above the running script).
    script_dir = os.path.dirname(os.path.abspath(__file__))
    current = script_dir
    for _ in range(6):
        current = os.path.dirname(current)
        if not current or current == os.path.dirname(current):
            break
        candidates.append(os.path.join(current, LOOKUP_FILE))
    for path in candidates:
        if os.path.isfile(path):
            return path
    return ""

def fetch_conversion_rate(product_code, pages, ruling, lamination_type):
    try:
        p_code = str(product_code).strip()
        if len(p_code) < 3 or p_code[2] not in ("3", "5"):
            return ""
        digit = int(p_code[2])
        target_page = _lookup_number(pages)
        if target_page is None:
            return ""
        target_ruling = _normalise_lookup_text(ruling)
        target_finish = _normalise_lookup_text(lamination_type)
        lookup_path = _find_conversion_lookup_file()
        if not lookup_path:
            return ""

        wb = openpyxl.load_workbook(lookup_path, data_only=True)
        for ws in wb.worksheets:
            header_row = None
            ruling_col_idx = None
            for row in range(1, min(ws.max_row, 10) + 1):
                for col in range(1, ws.max_column + 1):
                    header = _normalise_lookup_text(ws.cell(row=row, column=col).value)
                    if header == target_ruling:
                        header_row = row
                        ruling_col_idx = col
                        break
                if ruling_col_idx:
                    break
            if not ruling_col_idx:
                continue
            for row in range(header_row + 1, ws.max_row + 1):
                r_digit = _lookup_number(ws.cell(row=row, column=1).value)
                r_finish = _normalise_lookup_text(ws.cell(row=row, column=2).value)
                r_page = _lookup_number(ws.cell(row=row, column=3).value)
                if r_digit == digit and r_finish == target_finish and r_page == target_page:
                    val = ws.cell(row=row, column=ruling_col_idx).value
                    if val is None or str(val).strip() == "":
                        return ""
                    try:
                        return f"{float(val):.2f}"
                    except (ValueError, TypeError):
                        return str(val).strip()
        return ""
    except Exception:
        return ""

def calculate_po_request_full_row(row):
    p_code = str(row.get("product_code", "") or "").strip()
    mrp = row.get("mrp", "")
    order_qty = row.get("order_qty", "")
    prod_qty = row.get("produced_qty", "")
    p_rate = row.get("paper_rate", "")
    p_tax = row.get("tax_paper", "")
    p_cost = row.get("current_paper_cost", "")
    p_mill = row.get("mill_paper", "")
    b_rate = row.get("board_rate", "")
    b_tax = row.get("tax_board", "")
    b_cost = row.get("current_board_cost", "")
    b_mill = row.get("mill_board", "")
    lamination = row.get("lamination_type", "")

    try:
        page = int(p_code[-3:]) if len(p_code) >= 3 else 0
    except (ValueError, TypeError):
        page = 0

    fixed_120_codes = {
        "1132101132", "1132201140", "1132103132",
        "1132203140", "1132205140", "1132208140",
    }
    col6 = 120 if p_code in fixed_120_codes else page

    ruling_map = {"1": "Unruled", "3": "Single Line", "5": "Four Line", "8": "Med. Square"}
    last_4_first = p_code[-4:-3] if len(p_code) >= 4 else ""
    ruling = ruling_map.get(last_4_first, "")
    third_char = p_code[2:3] if len(p_code) >= 3 else ""

    if third_char == "5":
        inner_qty = 12 if col6 < 121 else 6
        case_qty = {52:192, 76:144, 100:120, 140:84, 176:72, 240:48}.get(col6, 0)
        board_gsm = 200 if page > 175 else 190
        sheet_width = 45.5
        sheet_length = 91
        paper_gsm = 57
        reel_width = 90
        cutoff = 43
    elif third_char == "3":
        inner_qty = 12 if col6 < 140 else 6
        case_qty = {56:288, 92:192, 120:144, 164:96, 176:96}.get(col6, 0)
        board_gsm = 190
        sheet_width = 98
        sheet_length = 77
        paper_gsm = 57
        reel_width = 97
        cutoff = 37
        if page <= 56:
            paper_gsm = 54
    else:
        inner_qty = 0
        case_qty = 0
        board_gsm = sheet_width = sheet_length = paper_gsm = reel_width = cutoff = ""

    paper_wt = ""
    board_wt = ""
    if third_char in ("3", "5"):
        base = (paper_gsm * reel_width * cutoff) / 20000 / 493.75
        paper_wt = base * (col6 / (16 if third_char == "3" else 12))
        base_board = (board_gsm * sheet_width * sheet_length) / 20000 / 487.5
        board_wt = base_board / (8 if third_char == "3" else 3)

    try:
        per_book_paper = float(paper_wt) * float(p_cost) if paper_wt and p_cost not in ("", None) else ""
        if per_book_paper != "":
            per_book_paper = f"{per_book_paper:.2f}"
    except (ValueError, TypeError):
        per_book_paper = ""

    try:
        per_book_board_base = float(board_wt) * float(b_cost) if board_wt and b_cost not in ("", None) else ""
        per_book_board = (per_book_board_base + per_book_board_base * 0.03) if per_book_board_base != "" else ""
        if per_book_board != "":
            per_book_board = f"{per_book_board:.2f}"
    except (ValueError, TypeError):
        per_book_board = ""

    conversion = fetch_conversion_rate(p_code, col6, ruling, lamination)

    return [
        row.get("po_date", ""), "PINNING",
        "24x18" if third_char == "3" else "29.7x21", "SOFT", p_code,
        col6, ruling, mrp, inner_qty, case_qty, order_qty, prod_qty,
        board_gsm, sheet_width, sheet_length, paper_gsm, reel_width, cutoff,
        f"{paper_wt:.6f}" if isinstance(paper_wt, (int, float)) else paper_wt,
        f"{board_wt:.6f}" if isinstance(board_wt, (int, float)) else board_wt,
        p_rate, p_tax, p_cost, p_mill, b_rate, b_tax, b_cost, b_mill,
        per_book_paper, per_book_board, conversion, lamination,
        ""
    ]

def get_po_release_rate(conn, product_code, order_month=None):
    """
    Auto-calculate PO Release Rate from the latest matching PO Request.
    Formula: PO Request column 29 + column 30 + column 31
             = Per Book Paper Cost + Per Book Board Cost + Conversion Rate.
    """
    code = str(product_code or "").strip()
    if len(code) != 10 or not code.isdigit():
        return 0.0

    try:
        if order_month:
            request_df = pd.read_sql_query(
                "SELECT * FROM po_request_entries "
                "WHERE product_code = ? AND po_month = ? "
                "ORDER BY id DESC LIMIT 1",
                conn,
                params=(code, order_month),
            )
        else:
            request_df = pd.read_sql_query(
                "SELECT * FROM po_request_entries "
                "WHERE product_code = ? ORDER BY id DESC LIMIT 1",
                conn,
                params=(code,),
            )

        # If the selected Order Month has no PO Request, use the latest
        # request for the same Product Code as a fallback.
        if request_df.empty and order_month:
            request_df = pd.read_sql_query(
                "SELECT * FROM po_request_entries "
                "WHERE product_code = ? ORDER BY id DESC LIMIT 1",
                conn,
                params=(code,),
            )

        if request_df.empty:
            return 0.0

        report_row = calculate_po_request_full_row(request_df.iloc[0].to_dict())

        # Report columns are 1-based: 29 = Per Book Paper Cost,
        # 30 = Per Book Board Cost, 31 = Conversion Rate.
        col_29 = float(report_row[28] or 0)
        col_30 = float(report_row[29] or 0)
        col_31 = float(report_row[30] or 0)

        return round(col_29 + col_30 + col_31, 2)
    except (ValueError, TypeError, sqlite3.Error, Exception):
        return 0.0

def calculate_purchase_order_report(conn):
    orders = pd.read_sql_query("SELECT * FROM order_entries ORDER BY id", conn)
    rows = []

    for _, r in orders.iterrows():
        code = str(r.get("product_code", "") or "").strip()
        page, ruling, third = parse_product_code(code)
        order_qty = float(pd.to_numeric(r.get("order_qty", 0), errors="coerce") or 0)

        production_qty = 0.0
        if code:
            prod = pd.read_sql_query(
                "SELECT ok_notebook FROM production_form_entries WHERE product_code = ?",
                conn, params=(code,)
            )
            if not prod.empty:
                production_qty = float(
                    pd.to_numeric(prod["ok_notebook"], errors="coerce").fillna(0).sum()
                )

        net_order = max(order_qty - production_qty, 0)

        col6 = page or 0
        if code in {"1132101132", "1132201140", "1132103132", "1132203140", "1132205140", "1132208140"}:
            col6 = 120

        if third == "5":
            paper_size, board_size = "29.7x21", "91x91x190"
            paper_gsm, reel_w, cutoff = 57, 90, 43
            board_gsm, sheet_w, sheet_l = (200 if page and page > 175 else 190), 45.5, 91
            paper_wt = (paper_gsm * reel_w * cutoff) / 20000 / 500 * (col6 / 12)
            board_wt = ((board_gsm * sheet_w * sheet_l) / 20000 / 487.5) / 3
            stock_paper_size = "90x57x43"
        elif third == "3":
            paper_size, board_size = "24x18", "77x98x190"
            paper_gsm = 54 if page is not None and page <= 56 else 57
            reel_w, cutoff = 97, 37
            board_gsm, sheet_w, sheet_l = 190, 98, 77
            paper_wt = (paper_gsm * reel_w * cutoff) / 20000 / 500 * (col6 / 16)
            board_wt = ((board_gsm * sheet_w * sheet_l) / 20000 / 487.5) / 8
            stock_paper_size = "97x37x57"
        else:
            paper_size = board_size = stock_paper_size = ""
            paper_wt = board_wt = 0.0

        required_paper = order_qty * paper_wt
        required_board = order_qty * board_wt

        paper_stock = 0.0
        if stock_paper_size:
            pdf = pd.read_sql_query(
                "SELECT opening_reel,inward_qty,process_qty,wastage "
                "FROM paper_detail_entries WHERE paper_size = ?",
                conn, params=(stock_paper_size,)
            )
            if not pdf.empty:
                for c in ("opening_reel", "inward_qty", "process_qty", "wastage"):
                    pdf[c] = pd.to_numeric(pdf[c], errors="coerce").fillna(0)
                paper_stock = float((pdf.opening_reel + pdf.inward_qty - pdf.process_qty - pdf.wastage).sum())

        board_stock = 0.0
        if board_size:
            bdf = pd.read_sql_query(
                "SELECT opening_stock,inward_qty,out_for_printing,wastage,printed_board_received,consumption "
                "FROM board_detail_entries WHERE board_size = ?",
                conn, params=(board_size,)
            )
            if not bdf.empty:
                for c in ("opening_stock", "inward_qty", "out_for_printing", "wastage", "printed_board_received", "consumption"):
                    bdf[c] = pd.to_numeric(bdf[c], errors="coerce").fillna(0)
                board_stock = float((bdf.opening_stock + bdf.inward_qty + bdf.printed_board_received - bdf.out_for_printing - bdf.wastage - bdf.consumption).sum())

        rows.append([
            r.get("order_month", ""),
            code,
            ruling,
            page or "",
            order_qty,
            net_order,
            paper_size,
            round(required_paper, 3),
            board_size,
            round(required_board, 3),
            round(paper_stock, 3),
            round(board_stock, 3),
            round(max(required_paper - paper_stock, 0), 3),
            round(max(required_board - board_stock, 0), 3)
        ])

    return pd.DataFrame(rows, columns=[
        "Order Month", "Product Code", "Ruling", "No. of Page",
        "Order Quantity", "Net Order", "Paper Size", "Required Paper",
        "Board Size", "Required Board", "Current Paper Stock",
        "Current Board Stock", "Net Paper Required", "Net Board Required"
    ])

# ==============================================================================
# STREAMLIT INTERFACE CONFIGURATION
# ==============================================================================
# GLOBAL FINANCIAL YEAR CONTROL (F2)
# ==============================================================================
def _financial_year_label(d=None):
    d = d or datetime.date.today()
    start_year = d.year if d.month >= 4 else d.year - 1
    return f"{start_year}-{str(start_year + 1)[-2:]}"

def _financial_year_options():
    current_start = datetime.date.today().year if datetime.date.today().month >= 4 else datetime.date.today().year - 1
    return [f"{y}-{str(y + 1)[-2:]}" for y in range(current_start - 5, current_start + 6)]

def render_financial_year_control():
    """Common Financial Year selector shown in all main working modules."""
    options = _financial_year_options()
    if st.session_state.financial_year not in options:
        options.append(st.session_state.financial_year)
        options = sorted(options, key=lambda x: int(x.split('-')[0]))

    fy_left, fy_right = st.columns([1, 5])
    with fy_left:
        selected_fy = st.selectbox(
            "Financial Year (F2)",
            options,
            index=options.index(st.session_state.financial_year),
            key="global_financial_year_selector"
        )
    with fy_right:
        st.markdown(
            f"<div style='padding-top:28px; font-size:13px; font-weight:700; color:#2C3E50;'>"
            f"Current Financial Year: <span style='color:#1E88E5;'>{selected_fy}</span> &nbsp; | &nbsp; F2: Change Financial Year</div>",
            unsafe_allow_html=True
        )
    st.session_state.financial_year = selected_fy

    # Keyboard shortcut: pressing F2 focuses/opens the common Financial Year selector.
    components.html("""
    <script>
    (function(){
      try {
        const p = window.parent;
        if (p.__subhFyF2Installed) return;
        p.__subhFyF2Installed = true;
        p.document.addEventListener('keydown', function(e){
          if (e.key === 'F2') {
            e.preventDefault();
            const labels = Array.from(p.document.querySelectorAll('label'));
            const saleDateLab = labels.find(x => (x.innerText || '').includes('Sale Date (F2)'));
            if (saleDateLab) return;
            const lab = labels.find(x => (x.innerText || '').includes('Financial Year (F2)'));
            if (lab) {
              const root = lab.parentElement && lab.parentElement.parentElement;
              const btn = root && root.querySelector('button');
              if (btn) { btn.focus(); btn.click(); }
            }
          }
        }, true);
      } catch(err) {}
    })();
    </script>
    """, height=0, width=0)
# ==============================================================================
# REPORT PREVIEW COMPONENT
# ==============================================================================
def _report_preview_value(value):
    """Format a single report cell for the print-style preview."""
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass
    if isinstance(value, float):
        if value.is_integer():
            return f"{value:,.0f}"
        return f"{value:,.2f}"
    if isinstance(value, int):
        return f"{value:,}"
    return str(value)


def show_report_preview(df, title, subtitle="", preview_key="report_preview",
                        company="SUBH PAPER COMPANY", opening=None, closing=None,
                        max_rows=250):
    """Show a print-style report preview based on the supplied ledger screenshot."""
    if df is None:
        return
    if not isinstance(df, pd.DataFrame):
        df = pd.DataFrame(df)

    state_key = f"_preview_visible_{preview_key}"
    if state_key not in st.session_state:
        st.session_state[state_key] = False

    if st.button("👁 Preview Report", key=f"btn_{preview_key}", use_container_width=False):
        st.session_state[state_key] = not st.session_state[state_key]

    if not st.session_state[state_key]:
        return

    preview_df = df.copy()
    truncated = False
    if len(preview_df) > max_rows:
        preview_df = preview_df.head(max_rows)
        truncated = True

    for col in preview_df.columns:
        if pd.api.types.is_numeric_dtype(preview_df[col]):
            preview_df[col] = preview_df[col].map(_report_preview_value)
        else:
            preview_df[col] = preview_df[col].map(lambda x: html.escape(_report_preview_value(x)))

    table_html = preview_df.to_html(
        index=False, border=0, classes="spc-preview-table", escape=False
    )

    balance_html = ""
    if opening is not None or closing is not None:
        parts = []
        if opening is not None:
            parts.append(f"<span><b>Opening Balance:</b> {html.escape(str(opening))}</span>")
        if closing is not None:
            parts.append(f"<span><b>Closing Balance:</b> {html.escape(str(closing))}</span>")
        balance_html = f"<div class='spc-preview-footer'>{''.join(parts)}</div>"

    note = "<div class='spc-preview-note'>Preview showing first 250 rows.</div>" if truncated else ""
    subtitle_html = f"<div class='spc-preview-subtitle'>{html.escape(str(subtitle))}</div>" if subtitle else ""

    preview_html = f"""
    <div class='spc-preview-page'>
      <div class='spc-preview-company'>{html.escape(str(company))}</div>
      <div class='spc-preview-address'>SUBH PAPER COMPANY • MANAGEMENT REPORT</div>
      <div class='spc-preview-title'>{html.escape(str(title))}</div>
      {subtitle_html}
      <div class='spc-preview-divider'></div>
      <div class='spc-preview-table-wrap'>{table_html}</div>
      {balance_html}
      {note}
      <div class='spc-preview-page-footer'>Report Preview • Generated from current ERP data</div>
    </div>
    """
    st.markdown(preview_html, unsafe_allow_html=True)
    if st.button("✖ Close Preview", key=f"close_{preview_key}"):
        st.session_state[state_key] = False
        st.rerun()

# ==============================================================================
# PAGE 2: MAIN WORKSPACE HUB
# ==============================================================================
# ==============================================================================
# PAGE 3: DASHBOARD
# ==============================================================================

def _safe_num(value):
    try:
        if value is None or value == "":
            return 0.0
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _dashboard_month_short(value):
    """Display order month compactly, e.g. September 2026 -> Sep. 2026."""
    text = str(value or "").strip()
    if not text:
        return ""
    month_map = {
        "january": "Jan.", "february": "Feb.", "march": "Mar.",
        "april": "Apr.", "may": "May", "june": "Jun.",
        "july": "Jul.", "august": "Aug.", "september": "Sep.",
        "october": "Oct.", "november": "Nov.", "december": "Dec.",
    }
    parts = text.replace("-", " ").replace("/", " ").split()
    if len(parts) >= 2 and parts[0].lower() in month_map:
        return f"{month_map[parts[0].lower()]} {parts[-1]}"
    return text

def _dashboard_paper_status(conn):
    sizes = ["97x37x57", "90x43x57", "97x37x54"]
    result = []
    for size in sizes:
        df = pd.read_sql_query(
            """SELECT opening_reel, inward_qty, process_qty, wastage, wip_paper
               FROM paper_detail_entries WHERE paper_size = ?""",
            conn, params=(size,)
        )
        if df.empty:
            reel = wip = total = 0.0
        else:
            for col in ("opening_reel", "inward_qty", "process_qty", "wastage", "wip_paper"):
                df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)
            reel = float((df["opening_reel"] + df["inward_qty"] - df["process_qty"] - df["wastage"]).sum())
            wip = float(df["wip_paper"].sum())
            total = reel + wip
        result.append((size, reel, wip, total))
    return result

def _dashboard_board_status(conn):
    sizes = ["45.5x91x190", "45.5x98x200", "91x91x190", "77x98x190"]
    display_names = {
        "Our Godown": "Our Godown",
        "Fortune Graphics": "Fortune",
        "Rich": "Rich Printer",
        "Rich Printer": "Rich Printer",
        "Vishruta Packaging": "Vishruta",
        "Salasar": "Salasar",
    }

    # Fetch actual parties from Board Detail. A party is shown only when it has
    # positive material in at least one board size.
    try:
        party_df = pd.read_sql_query(
            "SELECT DISTINCT party_name FROM board_detail_entries "
            "WHERE COALESCE(TRIM(party_name), '') <> '' ORDER BY party_name",
            conn,
        )
        party_order = party_df["party_name"].astype(str).tolist()
    except Exception:
        party_order = []

    raw_balances = {}
    wip_balances = {}

    for party in party_order:
        raw_values = []
        wip_values = []
        has_material = False

        for size in sizes:
            df = pd.read_sql_query(
                """SELECT opening_stock, inward_qty, out_for_printing, wastage,
                          printed_board_received, consumption
                   FROM board_detail_entries
                   WHERE board_size = ? AND party_name = ?""",
                conn,
                params=(size, party),
            )

            if df.empty:
                raw_value = 0.0
                wip_value = 0.0
            else:
                for col in (
                    "opening_stock", "inward_qty", "out_for_printing", "wastage",
                    "printed_board_received", "consumption",
                ):
                    df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)

                raw_value = float((
                    df["opening_stock"] + df["inward_qty"]
                    - df["out_for_printing"] - df["wastage"] - df["consumption"]
                ).sum())
                wip_value = float(df["printed_board_received"].sum())

            raw_value = max(raw_value, 0.0)
            wip_value = max(wip_value, 0.0)
            raw_values.append(raw_value)
            wip_values.append(wip_value)

            if raw_value > 0 or wip_value > 0:
                has_material = True

        if has_material:
            display_party = display_names.get(party, party)
            raw_balances[display_party] = raw_values
            wip_balances[display_party] = wip_values

    return sizes, raw_balances, wip_balances

def _dashboard_order_status(conn):
    orders = pd.read_sql_query(
        """SELECT order_month, product_code, ruling_type, no_of_page, order_qty
           FROM order_entries ORDER BY id""", conn
    )
    # One dashboard row per Purchase Order Month + Product Code (Barcode).
    grouped = {}
    for _, order in orders.iterrows():
        code = str(order.get("product_code", "") or "").strip()
        if not code:
            continue
        month = str(order.get("order_month", "") or "")
        key = (month, code)
        if key not in grouped:
            ruling = str(order.get("ruling_type", "") or "")
            if not ruling:
                _, ruling, _ = parse_product_code(code)
            grouped[key] = {
                "Order Month": _dashboard_month_short(month),
                "Barcode": code,
                "Ruling": ruling,
                "Order": 0.0,
            }
        grouped[key]["Order"] += _safe_num(order.get("order_qty"))

    rows = []
    for item in grouped.values():
        code = item["Barcode"]
        order_qty = item["Order"]

        prod_df = pd.read_sql_query(
            """SELECT ok_notebook FROM production_form_entries
               WHERE product_code = ?""", conn, params=(code,)
        )
        production = float(pd.to_numeric(prod_df["ok_notebook"], errors="coerce").fillna(0).sum()) if not prod_df.empty else 0.0

        des_df = pd.read_sql_query(
            """SELECT despatch_qty FROM despatch_form_entries
               WHERE product_code = ?""", conn, params=(code,)
        )
        despatched = float(pd.to_numeric(des_df["despatch_qty"], errors="coerce").fillna(0).sum()) if not des_df.empty else 0.0

        due_prod = max(order_qty - production, 0.0)
        due_desp = max(production - despatched, 0.0)
        cl_stock = max(production - despatched, 0.0)

        rows.append({
            "Order Month": item["Order Month"],
            "Barcode": code,
            "Ruling": item["Ruling"],
            "Order": order_qty,
            "Old": "-",
            "Production": production,
            "Due for Production": due_prod,
            "Despatched": despatched,
            "Due for Despatched": due_desp,
            "Cl. Stock": cl_stock,
            "Remarks": ""
        })
    return pd.DataFrame(rows)


# ==============================================================================
# EXPLICIT EXPORT LIST (star-import friendly; includes underscore helpers)
# ==============================================================================
__all__ = [n for n in globals() if not n.startswith('__')]
